#!/usr/bin/env python3
"""One-off ingest: NJ DOL WALL (Workplace Accountability in Labor List) -> wage_theft_records.

Source: https://www.nj.gov/labor/ea/assets/PDFs/Wall_Dataset.xlsx (verified HTTP 200,
valid xlsx; requires `pip install openpyxl`).
Actual workbook layout (verified 2026-09-14):
  sheet "WALL_09_4_2026": Name of Employer/DBA | Employer's Principal Address |
      County | Industry Type | Violation Type | Total Liability Owed Under Final
      Judgment / O... | Date Posted on the WALL | Multiple Addresses
  sheet "Add.Locations": Name of Employer/DBA | Address of Employer | Zip Code |
      County | State | Multiple Addresses (extra locations, not separately upserted)

Row mapping (wage_theft_records columns per schema.sql):
  case_id 'WALL-<spreadsheet row number>', source_agency 'NJ_DOL_WALL',
  provenance_type 'VERIFIED_PUBLIC_ACTION', status 'LIABILITY_OUTSTANDING',
  Total Liability -> settlement_amount, Date Posted -> findings_date,
  Industry Type -> industry_description, Violation Type -> violation_type.

Emits chunked .sql files (<=5000 upserts each) runnable by
`wrangler d1 execute <db> --remote --file`. Idempotent: INSERT ... ON CONFLICT(case_id) DO UPDATE.
(`--remote` required: without it wrangler writes to the local simulator.)
"""

import argparse
import datetime as dt
import os
import re
import sys
import urllib.request

WALL_URL = "https://www.nj.gov/labor/ea/assets/PDFs/Wall_Dataset.xlsx"
SOURCE_AGENCY = "NJ_DOL_WALL"
SOURCE_DOCKET_URL = WALL_URL  # dataset URL is the verified public source locator
CHUNK_SIZE = 5000

ADDR_RE = re.compile(
    r"^(?P<street>.*?),\s*(?P<city>[^,]+?),\s*(?P<state>[A-Z]{2})\s*(?P<zip>\d{5}(?:-\d{4})?)?\s*$"
)
DBA_RE = re.compile(r"\bd\s*/\s*b\s*/\s*a\s+(?P<trade>.+)$", re.IGNORECASE)


def sql_lit(v):
    if v is None or v == "":
        return "NULL"
    if isinstance(v, float):
        return repr(float(v))
    if isinstance(v, int):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def parse_address(raw):
    raw = (raw or "").strip().replace("\n", " ")
    m = ADDR_RE.match(raw)
    if not m:
        return raw, "", "NJ", ""
    return (
        m.group("street").strip(),
        m.group("city").strip(),
        m.group("state").strip(),
        (m.group("zip") or "").strip(),
    )


def split_dba(name):
    name = (name or "").strip().replace("\n", " ")
    m = DBA_RE.search(name)
    if m:
        legal = DBA_RE.sub("", name).strip().rstrip(",")
        return legal or name, m.group("trade").strip()
    return name, ""


def iso_date(v):
    if v is None or v == "":
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10]
    s = str(v).strip()[:10]
    try:
        dt.date.fromisoformat(s)
        return s
    except ValueError:
        return ""


def download(url, dest):
    req = urllib.request.Request(url, headers={"User-Agent": "civic-registry-ingest/1.0"})
    with urllib.request.urlopen(req, timeout=60) as r, open(dest, "wb") as f:
        if r.status != 200:
            raise RuntimeError(f"download failed: HTTP {r.status} for {url}")
        f.write(r.read())
    print(f"downloaded {url} -> {dest} ({os.path.getsize(dest)} bytes)")


WAGE_COLS = [
    "case_id", "source_agency", "respondent_legal_name", "trade_name", "address",
    "city", "state", "zip_code", "naics_code", "industry_description",
    "violation_type", "back_wages_recovered", "civil_penalties_assessed",
    "workers_affected", "repeat_violator", "status", "findings_date",
    "settlement_amount", "description", "provenance_type", "source_docket_url",
]


def upsert_sql(vals):
    cols = ", ".join(WAGE_COLS)
    vs = ", ".join(sql_lit(v) for v in vals)
    sets = ", ".join(f"{c}=excluded.{c}" for c in WAGE_COLS if c != "case_id")
    return (
        f"INSERT INTO wage_theft_records ({cols}) VALUES ({vs}) "
        f"ON CONFLICT(case_id) DO UPDATE SET {sets}, synced_at=CURRENT_TIMESTAMP;"
    )


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out-dir", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "out"))
    ap.add_argument("--limit", type=int, default=0, help="only process first N rows (testing)")
    ap.add_argument("--xlsx", default="", help="reuse a local xlsx instead of downloading")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        sys.exit("missing dependency: pip install openpyxl")

    xlsx_path = args.xlsx or os.path.join(args.out_dir, "Wall_Dataset.xlsx")
    os.makedirs(args.out_dir, exist_ok=True)
    if not args.xlsx:
        download(WALL_URL, xlsx_path)

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    print("sheets:", wb.sheetnames)
    ws = wb["WALL_09_4_2026"] if "WALL_09_4_2026" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header = [(c or "").strip() for c in rows[0]]
    print("columns:", header)

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
            for j, h in enumerate(header):  # headers may be truncated in previews; allow prefix match
                if h.startswith(n) or n.startswith(h):
                    return j
        raise RuntimeError(f"required column missing (looked for {names} in {header})")

    i_name = col("Name of Employer/DBA")
    i_addr = col("Employer's Principal Address")
    i_county = col("County")
    i_ind = col("Industry Type")
    i_viol = col("Violation Type")
    i_liab = col("Total Liability Owed Under Final Judgment / Order")
    i_posted = col("Date Posted on the WALL")

    statements = []
    for n, r in enumerate(rows[1:], start=2):  # n = spreadsheet row number
        if not r[i_name]:
            continue
        legal, trade = split_dba(r[i_name])
        street, city, state, zipc = parse_address(r[i_addr])
        try:
            liability = float(r[i_liab] or 0)
        except (TypeError, ValueError):
            liability = 0.0
        posted = iso_date(r[i_posted])
        county = (r[i_county] or "").strip() if r[i_county] else ""
        industry = (r[i_ind] or "").strip() if r[i_ind] else ""
        viol = (r[i_viol] or "").strip() if r[i_viol] else ""
        desc = (
            f"NJ DOL Workplace Accountability in Labor List (WALL): {viol}; "
            f"industry {industry}; county {county}; posted {posted}; "
            f"total liability ${liability:,.2f}."
        )
        statements.append(upsert_sql([
            f"WALL-{n}", SOURCE_AGENCY, legal.upper(), trade, street,
            city, state or "NJ", zipc, "", industry,
            viol, 0.0, 0.0,
            0, 0, "LIABILITY_OUTSTANDING", posted,
            liability, desc, "VERIFIED_PUBLIC_ACTION", SOURCE_DOCKET_URL,
        ]))
        if args.limit and len(statements) >= args.limit:
            break

    files = []
    for i in range(0, len(statements), CHUNK_SIZE):
        chunk = statements[i:i + CHUNK_SIZE]
        path = os.path.join(args.out_dir, f"nj-wall-{i // CHUNK_SIZE + 1:02d}.sql")
        with open(path, "w") as f:
            f.write("\n".join(chunk) + "\n")
        files.append(path)

    print(f"rowcount: {len(statements)} upserts -> {len(files)} file(s)")
    for p in files:
        print("wrote", p)


if __name__ == "__main__":
    main()
